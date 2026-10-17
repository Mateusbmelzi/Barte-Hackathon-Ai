"""Leitores: um por arquivo de origem. Leem os arquivos COMO VÊM (sem limpeza manual,
sem alterar nada no disco). Devolvem listas de dicts com Decimal para dinheiro."""
import csv
import email
import email.utils
import io
import re
import unicodedata
from collections import Counter, OrderedDict
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal
from email import policy
from pathlib import Path

import openpyxl

from . import config
from .dinheiro import dec, fmt_br


# ---------------------------------------------------------------- utilidades
def _sem_acento(s):
    s = unicodedata.normalize("NFKD", str(s))
    return "".join(c for c in s if not unicodedata.combining(c)).lower().strip()


def _data_br(s):
    s = (s or "").strip()
    return datetime.strptime(s, "%d/%m/%Y").date() if s else None


def _so_digitos(s, tamanho=14):
    if s is None:
        return None
    if isinstance(s, float):
        s = int(s)
    d = re.sub(r"\D", "", str(s))
    return d.zfill(tamanho) if d else None  # CNPJ numérico perde zero à esquerda (doc §2.1)


def _decodificar(raw, preferidos):
    for enc in preferidos:
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ValueError("encoding não reconhecido")


def _procurar(pasta):
    """Procura cada fonte. Devolve (achados, ausentes, ambiguos, erros) com mensagens prontas:
    ausente = 0 arquivo; ambíguo = mais de 1 numa fonte que aceita só 1."""
    pasta = Path(pasta)
    achados, ausentes, ambiguos, erros = {}, [], [], []
    todos = [p for p in pasta.iterdir() if p.is_file()]
    for chave, padrao in config.PADROES_ARQUIVOS.items():
        rx = re.compile(_glob_para_regex(padrao), re.IGNORECASE)
        m = sorted((p for p in todos if rx.fullmatch(p.name)), key=lambda p: p.name.lower())
        if chave in config.FONTES_MULTIPLAS:
            if not m:
                ausentes.append(chave)
                erros.append(f"{chave}: esperado 1 ou mais arquivos '{padrao}' em {pasta}, achei nenhum")
            else:
                achados[chave] = m
        elif len(m) != 1:
            (ambiguos if m else ausentes).append(chave)
            erros.append(f"{chave}: esperado 1 arquivo '{padrao}' em {pasta}, achei {[p.name for p in m]}")
        else:
            achados[chave] = m[0]
    return achados, ausentes, ambiguos, erros


def descobrir_arquivos(pasta):
    """Acha cada fonte por padrão de nome (config.PADROES_ARQUIVOS).
    Fontes comuns: exatamente 1 arquivo (Path). Fontes em config.FONTES_MULTIPLAS (Orbe): 1 ou mais,
    devolvidas como lista de Path em ordem de nome (a Orbe pode mandar arquivos de ajuste no mês).
    Estrito: falta ou repetição levanta FileNotFoundError (o fechamento usa descobrir_arquivos_parcial)."""
    achados, _, _, erros = _procurar(pasta)
    if erros:
        raise FileNotFoundError("\n".join(erros))
    return achados


def descobrir_arquivos_parcial(pasta):
    """Como descobrir_arquivos, mas fonte AUSENTE não é erro: devolve (achados, ausentes).
    Fonte repetida (2 arquivos onde se espera 1) continua sendo erro: não há como escolher sozinho."""
    achados, ausentes, ambiguos, erros = _procurar(pasta)
    if ambiguos:
        raise FileNotFoundError("\n".join(e for e in erros if e.split(":")[0] in ambiguos))
    return achados, ausentes


def achatar(arquivos):
    """Resultado de descobrir_arquivos -> lista simples de Path (as fontes múltiplas viram vários itens)."""
    return [p for v in arquivos.values() for p in (v if isinstance(v, list) else [v])]


def _glob_para_regex(padrao):
    return "".join(".*" if c == "*" else re.escape(c) for c in padrao)


# ---------------------------------------------------------------- 01 Orbe
def ler_orbe(caminho):
    texto = _decodificar(Path(caminho).read_bytes(), ["utf-8-sig", "cp1252"])
    linhas = []
    leitor = csv.DictReader(io.StringIO(texto, newline=""), delimiter=";")
    for n, r in enumerate(leitor, start=2):  # linha física (1 = cabeçalho)
        if not (r.get("id_transacao") or "").strip():
            continue
        parc = (r["qtd_parcelas"] or "").strip()
        linhas.append({
            "id": r["id_transacao"].strip(),
            "linha": n,
            "arquivo": Path(caminho).name,
            "data": _data_br(r["data_movimento"]),
            "ec": r["cod_estabelecimento"].strip(),
            "nome": r["nome_estabelecimento"].strip(),
            "tipo": r["tipo_lancamento"].strip().upper(),
            "modalidade": (r["modalidade"] or "").strip().upper(),
            "parcelas": int(parc) if parc else None,
            "bandeira": (r["bandeira"] or "").strip(),
            "bruto": dec(r["valor_bruto"]),
            "taxa_aplicada": dec(r["taxa_aplicada_pct"]),
            "valor_taxa": dec(r["valor_taxa"]),
            "liquido": dec(r["valor_liquido"]),
            "id_ref": (r["id_referencia"] or "").strip() or None,
            "data_orig": _data_br(r["data_venda_original"]),
        })
    return linhas


_CAMPOS_DUP = ("data", "ec", "tipo", "modalidade", "parcelas", "bruto", "taxa_aplicada", "valor_taxa", "liquido", "id_ref")


def ler_orbe_varios(caminhos):
    """Lê 1 ou N arquivos Orbe (em ordem de nome), concatena e resolve repetições de id_transacao.
    Cada linha leva 'arquivo' (nome) e 'linha' (linha física nele). Devolve (linhas, duplicadas).
    Repetição IDÊNTICA: conta uma vez (a 1ª vale, a repetida sai). Repetição com VALORES DIFERENTES (S52, critério
    nosso): vale a linha do arquivo de AJUSTE (a última em ordem de nome); a anterior sai, mas a linha que vale
    ocupa o lugar da anterior na ordem. Cada item de 'duplicadas' = campos da linha que SAIU + 'original' (a linha que
    vale) + 'identica' (mesmos valores?). Nunca contam em dobro."""
    if isinstance(caminhos, (str, Path)):
        caminhos = [caminhos]
    posicao, linhas, duplicadas = {}, [], []
    for c in sorted(map(Path, caminhos), key=lambda p: p.name.lower()):
        for l in ler_orbe(c):
            k = posicao.get(l["id"])
            if k is None:
                posicao[l["id"]] = len(linhas)
                linhas.append(l)
                continue
            o = linhas[k]
            identica = all(l[f] == o[f] for f in _CAMPOS_DUP)
            if identica:
                duplicadas.append({**l, "original": o, "identica": True})
            else:
                linhas[k] = l                                   # S52: o ajuste (arquivo posterior) substitui
                duplicadas.append({**o, "original": l, "identica": False})
    return linhas, duplicadas


# ---------------------------------------------------------------- 02 planilha de taxas
_COLS_TAXAS = {
    "lojista": "lojista", "cnpj": "cnpj", "ec orbe": "ec",
    "debito": "DEBITO", "credito a vista": "CREDITO A VISTA", "cred. 2-6x": "PARC_2_6",
    "cred. 7-12x": "PARC_7_12", "pix": "PIX", "antecipacao": "ANTECIPACAO",
    "spread remessa": "spread", "vigencia": "vigencia", "obs": "obs",
}
FAIXAS = ["DEBITO", "CREDITO A VISTA", "PARC_2_6", "PARC_7_12", "PIX", "ANTECIPACAO"]


def _pct(v):
    """Taxa da planilha -> Decimal em % (2.69). Aceita fração (0,0269) ou texto '2,69%'."""
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return None
    if isinstance(v, (int, float)):
        return dec(v) * 100
    m = re.fullmatch(r"\s*([\d.,]+)\s*%?\s*", str(v))
    if not m:
        raise ValueError(f"taxa não numérica: {v!r}")
    x = dec(m.group(1))
    return x if "%" in str(v) or x >= 1 else x * 100


def _spread(v):
    """Spread contratado como fração. Branco/'padrão' -> None (= padrão, CC03)."""
    if v is None or (isinstance(v, str) and _sem_acento(v) in ("", "padrao")):
        return None
    return _pct(v) / 100


def _vigencia(v):
    """Texto de vigência -> (inicio, fim, ok). H05: formato não reconhecido -> ok=False."""
    if isinstance(v, datetime):
        return v.date(), None, True
    s = _sem_acento(v or "")
    m = re.fullmatch(r"(?:a partir de\s*)?(\d{2}/\d{2}/\d{4})", s)
    if m:
        return _data_br(m.group(1)), None, True
    m = re.fullmatch(r"ate\s*(\d{2}/\d{2}/\d{4})", s)
    if m:
        return None, _data_br(m.group(1)), True
    return None, None, False


def ler_taxas(caminho):
    wb = openpyxl.load_workbook(caminho, data_only=True)
    abas = [n for n in wb.sheetnames if "vigente" in _sem_acento(n)]
    if len(abas) != 1:
        raise ValueError(f"planilha de taxas: esperada 1 aba 'vigentes', achei {wb.sheetnames}")
    ws = wb[abas[0]]
    cab_linha, mapa = None, {}
    for row in ws.iter_rows():
        if row[0].value and _sem_acento(row[0].value) == "lojista":
            cab_linha = row[0].row
            for c in row:
                if c.value is not None and _sem_acento(c.value) in _COLS_TAXAS:
                    mapa[_COLS_TAXAS[_sem_acento(c.value)]] = c.column
            break
    if cab_linha is None:
        raise ValueError("planilha de taxas: cabeçalho 'Lojista' não encontrado")
    contratos = []
    for r in range(cab_linha + 1, ws.max_row + 1):
        lojista = ws.cell(r, mapa["lojista"]).value
        cnpj = ws.cell(r, mapa["cnpj"]).value
        if not lojista or cnpj is None:
            continue  # linhas de nota de rodapé ("* spread em branco = padrão")
        val = lambda k: ws.cell(r, mapa[k]).value
        ec = val("ec")
        inicio, fim, ok = _vigencia(val("vigencia"))
        contratos.append({
            "linha": r,
            "lojista": str(lojista).strip(),
            "cnpj": _so_digitos(cnpj),
            "ec": str(int(ec)) if isinstance(ec, (int, float)) else (str(ec).strip() if ec else None),
            "taxas": {f: _pct(val(f)) for f in FAIXAS},
            "spread": _spread(val("spread")),
            "spread_bruto": val("spread"),
            "vigencia_texto": val("vigencia"),
            "inicio": inicio, "fim": fim, "vigencia_ok": ok,
            "obs": val("obs"),
            "riscado": bool(ws.cell(r, mapa["lojista"]).font.strike),  # linha riscada = inativa
        })
    return contratos


_COLS_ANTIGAS = {"lojista": "lojista", "cnpj": "cnpj", "debito": ["DEBITO"], "credito a vista": ["CREDITO A VISTA"],
                 "cred. parcelado": ["PARC_2_6", "PARC_7_12"], "pix": ["PIX"]}


def ler_taxas_antigas(caminho):
    """Abas que NÃO são a vigente (ex.: 'ANTIGO 2025 - nao usar'). Usadas só como FATO no dossiê de
    investigação (taxa aplicada coincide com tabela antiga?), nunca para calcular comissão.
    'Créd. parcelado' vale para as duas faixas de parcelas."""
    wb = openpyxl.load_workbook(caminho, data_only=True)
    out = []
    for nome in wb.sheetnames:
        if "vigente" in _sem_acento(nome):
            continue
        ws = wb[nome]
        mapa = None
        for row in ws.iter_rows():
            if mapa is None:
                if row[0].value and _sem_acento(row[0].value) == "lojista":
                    mapa = {_sem_acento(c.value): c.column for c in row
                            if c.value is not None and _sem_acento(c.value) in _COLS_ANTIGAS}
                continue
            val = {k: ws.cell(row[0].row, c).value for k, c in mapa.items()}
            if not val.get("lojista") or val.get("cnpj") is None:
                continue
            taxas = {}
            for k, faixas in _COLS_ANTIGAS.items():
                if isinstance(faixas, list) and k in val:
                    for f in faixas:
                        taxas[f] = _pct(val[k])
            out.append({"aba": nome, "linha": row[0].row, "lojista": str(val["lojista"]).strip(),
                        "cnpj": _so_digitos(val["cnpj"]), "taxas": taxas})
    return out


# ---------------------------------------------------------------- 03 e-mail Ponte (anexo em memória)
_COLS_PONTE = {
    "no operacao": "op", "data contratacao": "data", "data liquidacao": "data_liq",
    "cnpj cliente": "cnpj", "cliente": "cliente", "moeda": "moeda", "valor me": "valor_me",
    "taxa cliente": "taxa_cliente", "ptax ref.": "ptax", "valor r$": "valor_brl",
    "spread r$": "spread_brl", "spread %": "spread_pct", "spread estimado r$ (d-1)": "spread_estimado",
    "iof r$": "iof", "status compliance": "status", "observacao": "obs",
}
_NUM_PONTE = ["valor_me", "taxa_cliente", "ptax", "valor_brl", "spread_brl", "spread_pct", "spread_estimado", "iof"]


def _abrir_ponte(caminho):
    """Abre o e-mail e o anexo. Devolve (meta, ws, linha do cabeçalho ou None, mapa campo->coluna, cabeçalhos achados)."""
    with open(caminho, "rb") as f:
        msg = email.message_from_binary_file(f, policy=policy.default)
    anexo = None
    for parte in msg.iter_attachments():
        nome = parte.get_filename() or ""
        if nome.lower().endswith(".xlsx"):
            anexo = (nome, parte.get_payload(decode=True))
            break
    if anexo is None:
        raise ValueError("e-mail da Ponte sem anexo .xlsx")
    corpo = msg.get_body(("plain",))
    texto = corpo.get_content() if corpo else ""
    m = re.search(r"creditado em (\d{2}/\d{2}/\d{4})", texto)
    meta = {
        "anexo": anexo[0],
        "assunto": str(msg["subject"]),
        "data_repasse": _data_br(m.group(1)) if m else None,
    }
    wb = openpyxl.load_workbook(io.BytesIO(anexo[1]), data_only=True)
    ws = wb[wb.sheetnames[0]] if "Operações" not in wb.sheetnames else wb["Operações"]
    cab, mapa, achados = None, {}, []
    for row in ws.iter_rows():
        if row[0].value and _sem_acento(row[0].value).replace("º", "o") in ("no operacao", "n operacao"):
            cab = row[0].row
            for c in row:
                k = _sem_acento(c.value).replace("º", "o") if c.value else None
                if c.value is not None:
                    achados.append(str(c.value).strip())
                if k in _COLS_PONTE:
                    mapa[_COLS_PONTE[k]] = c.column
            break
    return meta, ws, cab, mapa, achados


def ler_ponte(caminho):
    """Lê o .eml, extrai o anexo .xlsx em memória e devolve (operações, metadados)."""
    meta, ws, cab, mapa, _ = _abrir_ponte(caminho)
    if cab is None:
        raise ValueError("relatório Ponte: cabeçalho 'Nº Operação' não encontrado")
    ops = []
    for r in range(cab + 1, ws.max_row + 1):
        d = ws.cell(r, mapa["data"]).value
        op = ws.cell(r, mapa["op"]).value
        if not op or not isinstance(d, datetime):
            continue  # linhas vazias e rodapé ("Total de operações:")
        o = {k: ws.cell(r, c).value for k, c in mapa.items()}
        o["op"] = str(op).strip()
        o["linha_excel"] = r
        o["data"] = d.date()
        o["data_liq"] = o["data_liq"].date() if isinstance(o.get("data_liq"), datetime) else None
        o["cnpj"] = _so_digitos(o["cnpj"])
        for k in _NUM_PONTE:
            o[k] = dec(o[k])
        o["obs"] = o.get("obs") or ""
        ops.append(o)
    return ops, meta


# ---------------------------------------------------------------- 04 plano de contas
def ler_plano(caminho):
    texto = _decodificar(Path(caminho).read_bytes(), ["utf-8-sig", "cp1252"])
    plano = {}
    for r in csv.DictReader(io.StringIO(texto, newline=""), delimiter=";"):
        cod = (r.get("COD_CONTA") or "").strip()
        if cod:
            plano[cod] = {
                "descricao": r["DESCRICAO"].strip(),
                "tipo": r["TIPO"].strip(),
                "aceita": r["ACEITA_LANCTO"].strip().upper() == "S",
            }
    return plano


# ---------------------------------------------------------------- 05a Aurora (CSV)
def ler_aurora(caminho):
    """ref_origem = 'L' + número da linha física do arquivo (1ª linha = L1)."""
    texto = _decodificar(Path(caminho).read_bytes(), ["utf-8-sig", "cp1252"])
    movs, saldo_ini, saldo_fim = [], None, None
    for n, bruta in enumerate(texto.split("\n"), start=1):
        campos = bruta.rstrip("\r").split(";")
        campos += [""] * (5 - len(campos))
        data_s, desc, doc, valor_s, saldo_s = (c.strip() for c in campos[:5])
        desc_n = _sem_acento(desc)
        if desc_n == "saldo anterior":
            saldo_ini = dec(saldo_s)
        elif desc_n == "saldo final":
            saldo_fim = dec(saldo_s)
        elif re.fullmatch(r"\d{2}/\d{2}/\d{4}", data_s) and valor_s:
            movs.append({
                "ref": f"L{n}", "linha": n, "data": _data_br(data_s), "descricao": desc,
                "documento": doc, "valor": dec(valor_s), "saldo": dec(saldo_s),
            })
    return movs, {"saldo_anterior": saldo_ini, "saldo_final": saldo_fim}


# ---------------------------------------------------------------- 05b Meridiano (OFX SGML)
def ler_meridiano(caminho):
    raw = Path(caminho).read_bytes()
    m = re.search(rb"CHARSET:(\d+)", raw)
    enc = f"cp{m.group(1).decode()}" if m else "cp1252"
    texto = raw.decode(enc)
    movs = []
    for bloco in re.findall(r"<STMTTRN>(.*?)</STMTTRN>", texto, re.S):
        t = dict(re.findall(r"<([A-Z0-9.]+)>([^\r\n<]*)", bloco))
        movs.append({
            "ref": t["FITID"].strip(), "data": datetime.strptime(t["DTPOSTED"][:8], "%Y%m%d").date(),
            "tipo": t.get("TRNTYPE", "").strip(), "valor": dec(t["TRNAMT"].strip()),
            "checknum": t.get("CHECKNUM", "").strip(), "descricao": t.get("MEMO", "").strip(),
        })
    lb = re.search(r"<LEDGERBAL>\s*<BALAMT>([^\r\n<]+)", texto)
    return movs, {"saldo_final": dec(lb.group(1).strip()) if lb else None}


def ler_tudo(pasta):
    arq = descobrir_arquivos(pasta)
    ponte, meta_ponte = ler_ponte(arq["ponte"])
    aurora, meta_aurora = ler_aurora(arq["aurora"])
    meridiano, meta_mer = ler_meridiano(arq["meridiano"])
    orbe, orbe_dup = ler_orbe_varios(arq["orbe"])
    arquivos = {k: v.name for k, v in arq.items() if k != "orbe"}
    arquivos["orbe"] = ", ".join(p.name for p in arq["orbe"])
    return {
        "arquivos": arquivos,
        "arquivos_orbe": [p.name for p in arq["orbe"]],
        "orbe": orbe, "orbe_duplicadas": orbe_dup,
        "contratos": ler_taxas(arq["taxas"]),
        "ponte": ponte, "meta_ponte": meta_ponte,
        "plano": ler_plano(arq["plano"]),
        "aurora": aurora, "meta_aurora": meta_aurora,
        "meridiano": meridiano, "meta_meridiano": meta_mer,
    }


# ---------------------------------------------------------------- leitura tolerante (fonte ausente / layout)
# Campos que cada leitor usa. Se faltar algum, a fonte é "com layout inesperado" (S53): o fechamento
# segue sem ela e avisa quais colunas faltam, em vez de quebrar.
_COLS_ORBE = ["id_transacao", "data_movimento", "cod_estabelecimento", "nome_estabelecimento", "tipo_lancamento",
              "modalidade", "qtd_parcelas", "bandeira", "valor_bruto", "taxa_aplicada_pct", "valor_taxa",
              "valor_liquido", "id_referencia", "data_venda_original"]
_COLS_PLANO = ["COD_CONTA", "DESCRICAO", "TIPO", "ACEITA_LANCTO"]
_COLS_AURORA = ["Data", "Lançamento", "Documento", "Valor (R$)", "Saldo (R$)"]
_TAGS_OFX = ["FITID", "DTPOSTED", "TRNAMT"]
_ROTULOS_TAXAS = {"lojista": "Lojista", "cnpj": "CNPJ", "ec orbe": "EC Orbe", "debito": "Débito",
                  "credito a vista": "Crédito à vista", "cred. 2-6x": "Créd. 2-6x", "cred. 7-12x": "Créd. 7-12x",
                  "pix": "PIX", "antecipacao": "Antecipação", "spread remessa": "Spread remessa",
                  "vigencia": "Vigência", "obs": "Obs"}


def _texto_arquivo(caminho):
    return _decodificar(Path(caminho).read_bytes(), ["utf-8-sig", "cp1252"])


def _cabecalho_csv(caminho):
    return [c.strip() for c in _texto_arquivo(caminho).split("\n", 1)[0].rstrip("\r").split(";")]


def _validar_orbe(caminho):
    cab = _cabecalho_csv(caminho)
    return [c for c in _COLS_ORBE if c not in cab], cab


def _validar_plano(caminho):
    cab = _cabecalho_csv(caminho)
    return [c for c in _COLS_PLANO if c not in cab], cab


def _validar_taxas(caminho):
    wb = openpyxl.load_workbook(caminho, data_only=True)
    abas = [n for n in wb.sheetnames if "vigente" in _sem_acento(n)]
    if len(abas) != 1:
        return ["aba 'vigentes'"], wb.sheetnames
    for row in wb[abas[0]].iter_rows():
        if row[0].value and _sem_acento(row[0].value) == "lojista":
            achadas = [str(c.value).strip() for c in row if c.value is not None]
            norm = {_sem_acento(a) for a in achadas}
            return [_ROTULOS_TAXAS[k] for k in _COLS_TAXAS if k not in norm], achadas
    return ["cabeçalho 'Lojista'"], [f"aba '{abas[0]}' sem linha de cabeçalho"]


def _validar_ponte(caminho):
    _, _, cab, mapa, achadas = _abrir_ponte(caminho)
    if cab is None:
        return ["cabeçalho 'Nº Operação'"], ["(cabeçalho não encontrado)"]
    nomes = {v: k for k, v in _COLS_PONTE.items()}
    return [nomes[c] for c in _COLS_PONTE.values() if c not in mapa], achadas


def _validar_aurora(caminho):
    linhas = [l.rstrip("\r").split(";") for l in _texto_arquivo(caminho).split("\n")]
    esperado = [_sem_acento(c) for c in _COLS_AURORA]
    cab = next((l for l in linhas if len(l) >= 5 and _sem_acento(l[0]) == "data"), None)
    if cab is None:
        achadas = [c.strip() for c in linhas[0]] if linhas else []
        return ["linha de cabeçalho '" + ";".join(_COLS_AURORA) + "'"], achadas
    achadas = [c.strip() for c in cab]
    faltam = [orig for orig, ex, got in zip(_COLS_AURORA, esperado, cab[:5]) if _sem_acento(got) != ex]
    if not faltam and not any(_sem_acento(l[1]) == "saldo anterior" for l in linhas if len(l) > 1):
        faltam = ["linha 'SALDO ANTERIOR'"]
    return faltam, achadas


def _validar_meridiano(caminho):
    texto = Path(caminho).read_bytes().decode("cp1252", errors="replace")
    blocos = re.findall(r"<STMTTRN>(.*?)</STMTTRN>", texto, re.S)
    if not blocos:
        return ["blocos <STMTTRN>"], ["(nenhum movimento encontrado)"]
    tags = set()
    for b in blocos:
        tags |= set(re.findall(r"<([A-Z0-9.]+)>", b))
    return [t for t in _TAGS_OFX if not all(f"<{t}>" in b for b in blocos)], sorted(tags)


_VALIDADORES = {"orbe": _validar_orbe, "taxas": _validar_taxas, "ponte": _validar_ponte, "plano": _validar_plano,
                "aurora": _validar_aurora, "meridiano": _validar_meridiano}


def validar_layout(chave, caminho):
    """None se o arquivo tem os campos que o leitor usa; senão {'faltantes': [...], 'encontradas': [...]}.
    Arquivo ilegível (não é o formato esperado) também é tratado como layout inesperado."""
    try:
        faltantes, achadas = _VALIDADORES[chave](caminho)
    except Exception as e:  # noqa: BLE001 - qualquer falha de abertura = layout inesperado, nunca stack trace
        return {"faltantes": [f"arquivo não pôde ser aberto como esperado ({type(e).__name__}: {e})"],
                "encontradas": []}
    return {"faltantes": faltantes, "encontradas": achadas} if faltantes else None


# ---------------------------------------------------------------- versões (2+ arquivos da mesma fonte)
def _nat(nome):
    """Ordem de nome 'natural' (v2 < v10); 'x (1).csv' (duplicata do navegador) vem depois de 'x.csv'."""
    n = re.sub(r"\s*\((\d+)\)(\.\w+)$", r"~\1\2", nome.lower())
    return [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", n)]


def _momento_eml(caminho):
    """Cabeçalho Date do e-mail, em UTC sem fuso."""
    with open(caminho, "rb") as f:
        msg = email.message_from_binary_file(f, policy=policy.default)
    dt = email.utils.parsedate_to_datetime(str(msg["Date"]))
    return dt.astimezone(timezone.utc).replace(tzinfo=None) if dt.tzinfo else dt


def _momento_ofx(caminho):
    """DTSERVER do OFX (ex.: 20261001063000[-3:BRT]), em UTC."""
    texto = Path(caminho).read_bytes().decode("cp1252", errors="replace")
    m = re.search(r"<DTSERVER>(\d{8})(\d{6})?(?:\[([+-]?\d+(?:\.\d+)?)(?::[A-Z]+)?\])?", texto)
    if not m:
        return None
    dt = datetime.strptime(m.group(1) + (m.group(2) or "000000"), "%Y%m%d%H%M%S")
    return dt - timedelta(hours=float(m.group(3))) if m.group(3) else dt


def _momento_aurora(caminho):
    """Fim do 'Período: dd/mm/aaaa a dd/mm/aaaa' do cabeçalho do CSV."""
    m = re.search(r"Per.odo:\s*\d{2}/\d{2}/\d{4}\s*a\s*(\d{2}/\d{2}/\d{4})", _texto_arquivo(caminho))
    return datetime.combine(_data_br(m.group(1)), time.min) if m else None


_MOMENTO_VERSAO = {"ponte": _momento_eml, "meridiano": _momento_ofx, "aurora": _momento_aurora}


def ordenar_versoes(chave, caminhos):
    """Da versão mais antiga para a mais nova: data/hora embutida (Ponte Date, OFX DTSERVER, Aurora fim do Período)
    e, no empate ou sem data, ordem de nome natural. Orbe/taxas/plano (sem data embutida): ordem de nome."""
    def k(p):
        try:
            dt = _MOMENTO_VERSAO[chave](p) if chave in _MOMENTO_VERSAO else None
        except Exception:  # noqa: BLE001 - arquivo ilegível: o validador de layout reclama depois
            dt = None
        return (dt or datetime.min, _nat(p.name))
    return sorted(caminhos, key=k)


def descobrir_versoes(pasta):
    """Como descobrir_arquivos_parcial, mas TODAS as fontes aceitam 1 ou mais arquivos (versões).
    Devolve (achados {fonte: [Path em ordem de versão, o último é o mais novo]}, ausentes)."""
    pasta = Path(pasta)
    todos = [p for p in pasta.iterdir() if p.is_file()]
    achados, ausentes = {}, []
    for chave, padrao in config.PADROES_ARQUIVOS.items():
        rx = re.compile(_glob_para_regex(padrao), re.IGNORECASE)
        m = [p for p in todos if rx.fullmatch(p.name)]
        if not m:
            ausentes.append(chave)
        elif chave == "orbe":
            achados[chave] = sorted(m, key=lambda p: p.name.lower())
        else:
            achados[chave] = ordenar_versoes(chave, m)
    return achados, ausentes


def _pl(n, sing, plur):
    return f"{n} {sing if n == 1 else plur}"


_CAMPOS_DINHEIRO = ("valor", "valor_brl", "spread_brl", "spread_estimado", "iof")


def _fmt_v(campo, v):
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    if v is None or v == "":
        return "(vazio)"
    if isinstance(v, Decimal) and campo in _CAMPOS_DINHEIRO:
        return fmt_br(v)
    return str(v)


def _difs(campos, antigo, novo):
    return "; ".join(f"{c}: {_fmt_v(c, antigo[c])} -> {_fmt_v(c, novo[c])}" for c in campos if antigo[c] != novo[c])


def _pend_versao(tipo, ref, valor, motivo, rastro):
    return {"tipo": tipo, "ref_origem": ref, "valor": valor, "motivo": motivo,
            "quem_decide": config.QUEM_HELENA, "rastro": rastro}


def _consolidar_chaveada(versoes, chave_de, campos):
    """Consolida versões por chave natural. versoes = [(arquivo, linhas)] da mais antiga para a mais nova.
    A ordem e o conteúdo vêm da versão mais nova; a linha que só existia em versão antiga é MANTIDA (a da versão mais
    recente que a continha). Devolve (linhas [(linha, arquivo)], chaves novas, conflitos, sumiram):
      conflitos = [(chave, linha nova, arquivo novo, [(arquivo antigo, linha antiga)])] (valem os da nova);
      sumiram   = [(chave, 1ª linha, arquivo)]."""
    n = len(versoes)
    grupos = []
    for _, linhas in versoes:
        g = OrderedDict()
        for l in linhas:
            g.setdefault(chave_de(l), []).append(l)
        grupos.append(g)
    nova = grupos[-1]
    out = [(l, versoes[-1][0]) for l in versoes[-1][1]]
    dono = {}
    for i in range(n - 1):
        for k in grupos[i]:
            if k not in nova:
                dono[k] = i                      # a versão mais recente (antes da nova) que ainda tinha a chave
    sumiram = []
    for i in range(n - 1):
        for k, ls in grupos[i].items():
            if dono.get(k) == i:
                out.extend((l, versoes[i][0]) for l in ls)
                sumiram.append((k, ls[0], versoes[i][0]))
    conflitos = []
    for k, ls in nova.items():
        antigos = [(versoes[i][0], grupos[i][k][0]) for i in range(n - 1)
                   if k in grupos[i] and any(grupos[i][k][0][c] != ls[0][c] for c in campos)]
        if antigos:
            conflitos.append((k, ls[0], versoes[-1][0], antigos))
    novas = [k for k in nova if not any(k in grupos[i] for i in range(n - 1))]
    return out, novas, conflitos, sumiram


def _chave_aurora(m):
    return (m["data"], m["descricao"], m["documento"], m["valor"])


def _contem(maior, menor):
    return all(maior[k] >= v for k, v in menor.items())


def _consolidar_aurora(versoes):
    """versoes = [(arquivo, movs, meta)] da mais antiga para a mais nova. Multiconjunto por
    (data, lançamento, documento, valor).
      - Alguma versão contém todas as outras (a mais nova, se houver empate) -> usa só ela, com as linhas físicas dela
        (ref 'Lnn' do edital). Se a escolhida não for a mais nova, o que a nova NÃO tem vira 'sumiram'.
      - Nenhuma contém as outras (sobreposição parcial) -> união pelo máximo de ocorrências, ref 'Lnn@arquivo'.
    Devolve dict: movs, meta, usados, novas (qtd), sumiram [mov], sobreposta."""
    n = len(versoes)
    cont = [Counter(_chave_aurora(m) for m in movs) for _, movs, _ in versoes]
    nova = cont[-1]
    antigo_max = Counter()
    for c in cont[:-1]:
        for k, v in c.items():
            antigo_max[k] = max(antigo_max[k], v)
    novas = sum(max(0, v - antigo_max[k]) for k, v in nova.items())
    base = next((i for i in range(n - 1, -1, -1) if all(_contem(cont[i], cont[j]) for j in range(n))), None)
    if base is not None:
        arq, movs, meta = versoes[base]
        saida = [{**m, "arquivo": arq} for m in movs]
        sumiram = []
        if base != n - 1:
            falta = {k: v - nova[k] for k, v in cont[base].items() if v > nova[k]}
            for m in reversed(saida):
                k = _chave_aurora(m)
                if falta.get(k, 0) > 0:
                    falta[k] -= 1
                    sumiram.append(m)
            sumiram.reverse()
        return {"movs": saida, "meta": meta, "usados": [arq], "novas": novas, "sumiram": sumiram, "sobreposta": False}
    arq_n, movs_n, _ = versoes[-1]
    saida = [{**m, "ref": f"{m['ref']}@{arq_n}", "arquivo": arq_n} for m in movs_n]
    tem = Counter(nova)
    sumiram = []
    for i in range(n - 2, -1, -1):
        arq, movs, _ = versoes[i]
        for m in movs:
            k = _chave_aurora(m)
            if tem[k] < antigo_max[k]:
                tem[k] += 1
                extra = {**m, "ref": f"{m['ref']}@{arq}", "arquivo": arq}
                saida.append(extra)
                sumiram.append(extra)
    saida.sort(key=lambda m: m["data"])          # estável: mantém a ordem do arquivo dentro do mesmo dia
    return {"movs": saida, "meta": {"saldo_anterior": None, "saldo_final": None},
            "usados": [a for a, _, _ in versoes], "novas": novas, "sumiram": sumiram, "sobreposta": True}


def _entrada_versao(fonte, arquivos, usados, novas=0, conflitos=0, sumiram=0, sobreposta=False):
    """Resumo de versões de uma fonte (r['versoes'][fonte]); 'texto' é a linha do resumo do CLI."""
    ignorados = [a for a in arquivos if a not in usados]
    rotulo = config.ROTULO_VERSAO[fonte]
    unid = config.UNIDADE_VERSAO.get(fonte)
    if fonte == "orbe":
        txt = f"{rotulo}: {len(arquivos)} arquivos — {' + '.join(usados)}"
    else:
        txt = (f"{rotulo}: {len(arquivos)} versões — " + ("usadas " if len(usados) > 1 else "usada ")
               + " + ".join(usados))
    if unid is None:
        txt += f" (ignorada {', '.join(ignorados)})" if ignorados else ""
    else:
        det = []
        if sobreposta:
            det.append("união por sobreposição")
        det.append(f"+{_pl(novas, f'{unid[0]} nov{unid[2]}', f'{unid[1]} nov{unid[2]}s')}" if novas
                   else f"sem {unid[1]} nov{unid[2]}s")
        if conflitos:
            det.append(_pl(conflitos, "conflito de valor", "conflitos de valor"))
        if sumiram:
            det.append(_pl(sumiram, "sumiu da nova versão", "sumiram da nova versão"))
        txt += " (" + "; ".join(det) + ")"
    return {"fonte": fonte, "rotulo": rotulo, "arquivos": list(arquivos), "usados": list(usados),
            "ignorados": ignorados, "novas": novas, "conflitos": conflitos, "sumiram": sumiram,
            "sobreposta": sobreposta, "texto": txt}


def _usados(out, nome_novo):
    return [nome_novo] + [a for a in dict.fromkeys(a for _, a in out) if a != nome_novo]


def _consolidar_ofx(d, lidos):
    """lidos = [(Path, (movs, meta))] em ordem de versão. FITID é a chave; vale a versão mais nova."""
    versoes = [(c.name, r[0]) for c, r in lidos]
    campos = ["data", "tipo", "valor", "checknum", "descricao"]
    out, novas, conflitos, sumiram = _consolidar_chaveada(versoes, lambda m: m["ref"], campos)
    movs = [{**m, "arquivo": arq} for m, arq in out]
    nome_novo = versoes[-1][0]
    pend = []
    for k, m, arq, antigos in conflitos:
        quais = " | ".join(f"{a}: {_difs(campos, ant, m)}" for a, ant in antigos)
        pend.append(_pend_versao(
            config.TIPO_VERSAO_CONFLITANTE, k, m["valor"],
            f"FITID {k} aparece com valores diferentes no extrato Meridiano. Vale a versão mais recente ({arq}): "
            f"{_fmt_v('valor', m['valor'])} em {_fmt_v('data', m['data'])}, '{m['descricao']}'. Versão anterior -> mais "
            f"recente: {quais}. Critério nosso: a mais recente vale; Helena decide qual está certa.",
            [(f"meridiano@{arq}", k, m["valor"], "versão usada (mais recente)")]))
    for k, m, arq in sumiram:
        pend.append(_pend_versao(
            config.TIPO_LANCAMENTO_SUMIU, k, m["valor"],
            f"O lançamento {k} ('{m['descricao']}', {_fmt_v('data', m['data'])}, {_fmt_v('valor', m['valor'])}) existia em "
            f"{arq} e não consta em {nome_novo}. Mantido no fechamento (nada some em silêncio); pode ser estorno ou "
            "correção do banco. Helena confirma se continua valendo.",
            [(f"meridiano@{arq}", k, m["valor"], "lançamento que sumiu da nova versão (mantido)")]))
    for m in movs:
        d["rastro_arquivos"][f"meridiano@{m['arquivo']}"] = m["arquivo"]
    d["versoes"]["meridiano"] = _entrada_versao("meridiano", [a for a, _ in versoes], _usados(out, nome_novo),
                                                len(novas), len(conflitos), len(sumiram))
    d["pend_versoes"] += pend
    return movs, lidos[-1][1][1]


def _consolidar_ponte(d, lidos):
    """lidos = [(Path, (ops, meta))]. Nº Operação é a chave entre arquivos; a repetição dentro do MESMO relatório fica
    como está (anomalia R24, tratada em remessa.calcular)."""
    versoes = [(c.name, r[0]) for c, r in lidos]
    metas = {c.name: r[1] for c, r in lidos}
    campos = ["data", "data_liq", "cnpj", "cliente", "moeda", "valor_me", "taxa_cliente", "ptax", "valor_brl",
              "spread_brl", "spread_pct", "spread_estimado", "iof", "status", "obs"]
    out, novas, conflitos, sumiram = _consolidar_chaveada(versoes, lambda o: o["op"], campos)
    ops = [{**o, "arquivo": arq} for o, arq in out]
    nome_novo = versoes[-1][0]
    d["origem_ponte"] = {o["op"]: o["arquivo"] for o in ops}
    pend = []
    for k, o, arq, antigos in conflitos:
        quais = " | ".join(f"{a}: {_difs(campos, ant, o)}" for a, ant in antigos)
        pend.append(_pend_versao(
            config.TIPO_VERSAO_CONFLITANTE, k, o["spread_brl"],
            f"A operação {k} ({o['cliente']}) aparece com valores diferentes em relatórios da Ponte. Vale a versão mais "
            f"recente ({arq}), Spread R$ {fmt_br(o['spread_brl'])}. Versão anterior -> mais recente: {quais}. "
            "Critério nosso: a mais recente vale; Helena decide qual está certa e confirma com a Ponte.",
            [(f"ponte@{arq}", k, o["spread_brl"], "versão usada (mais recente)")]))
    for k, o, arq in sumiram:
        pend.append(_pend_versao(
            config.TIPO_LANCAMENTO_SUMIU, k, o["spread_brl"],
            f"A operação {k} ({o['cliente']}, {_fmt_v('data', o['data'])}, Spread R$ {fmt_br(o['spread_brl'])}) existia em "
            f"{arq} e não consta em {nome_novo}. Mantida no fechamento (nada some em silêncio); pode ser cancelamento ou "
            "correção da Ponte. Helena confirma com a Ponte.",
            [(f"ponte@{arq}", k, o["spread_brl"], "operação que sumiu da nova versão (mantida)")]))
    for a in {a for _, a in out}:
        d["rastro_arquivos"][f"ponte@{a}"] = f"{a}#{metas[a]['anexo']}"
    d["versoes"]["ponte"] = _entrada_versao("ponte", [a for a, _ in versoes], _usados(out, nome_novo), len(novas),
                                            len(conflitos), len(sumiram))
    d["pend_versoes"] += pend
    return ops, metas[nome_novo]


def _consolidar_aurora_fonte(d, lidos):
    versoes = [(c.name, r[0], r[1]) for c, r in lidos]
    res = _consolidar_aurora(versoes)
    nome_novo = versoes[-1][0]
    pend = []
    if res["sobreposta"]:
        fecha, falhas = [], []
        for arq, movs, meta in versoes:
            si, sf = meta.get("saldo_anterior"), meta.get("saldo_final")
            mov = sum((m["valor"] for m in movs), Decimal("0"))
            if si is None or sf is None or si + mov == sf:
                fecha.append(f"{arq}: fecha")
            else:
                txt = f"saldo anterior {fmt_br(si)} + movimentos {fmt_br(mov)} != saldo final {fmt_br(sf)}"
                fecha.append(f"{arq}: NÃO fecha ({txt})")
                falhas.append(_pend_versao("EXTRATO_NAO_FECHA", f"AURORA ({arq})", None,
                                           f"{txt[0].upper()}{txt[1:]} na versão {arq}.", []))
        antigo_max = Counter()
        for _, movs, _ in versoes[:-1]:
            for k, v in Counter(_chave_aurora(m) for m in movs).items():
                antigo_max[k] = max(antigo_max[k], v)
        vistas, so_nova = Counter(), []
        for m in versoes[-1][1]:                       # linhas da mais recente que nenhuma versão anterior tem
            k = _chave_aurora(m)
            vistas[k] += 1
            if vistas[k] > antigo_max[k]:
                so_nova.append(m)
        cand = [f"'{x['descricao']}' de {_fmt_v('data', x['data'])}: {fmt_br(x['valor'])} em {x['arquivo']} -> "
                f"{fmt_br(n['valor'])} em {nome_novo}"
                for x in res["sumiram"] for n in so_nova
                if x["data"] == n["data"] and x["descricao"] == n["descricao"] and x["valor"] != n["valor"]]
        aviso_corr = (" Se a versão nova corrige um lançamento da antiga, remova a versão antiga na aba Fechamento "
                      "(botão 'remover desta pasta') e rode de novo; enquanto isso, os dois lançamentos estão contados."
                      + (" Candidatos (mesma data e descrição, valor diferente): " + "; ".join(cand) + "." if cand else ""))
        pend.append(_pend_versao(
            config.TIPO_VERSAO_SOBREPOSTA, "AURORA", None,
            f"As {len(versoes)} versões do extrato Aurora ({', '.join(a for a, _, _ in versoes)}) se sobrepõem sem que uma "
            "contenha a outra. Foi feita a UNIÃO (cada lançamento pelo máximo de ocorrências em qualquer versão, "
            "comparando data, lançamento, documento e valor); ref_origem fica 'Lnn@arquivo'. "
            f"{_pl(res['novas'], 'lançamento novo', 'lançamentos novos')} na mais recente ({nome_novo}); "
            f"{_pl(len(res['sumiram']), 'só existe', 'só existem')} em versão anterior (ver "
            f"{config.TIPO_LANCAMENTO_SUMIU}). Saldo por versão: {'; '.join(fecha)}. O saldo do conjunto unido não é "
            "conferido. Helena decide qual versão vale." + aviso_corr, []))
        pend += falhas
    for m in res["sumiram"]:
        pend.append(_pend_versao(
            config.TIPO_LANCAMENTO_SUMIU, m["ref"], m["valor"],
            f"O lançamento '{m['descricao']}' de {_fmt_v('data', m['data'])} (doc. {m['documento'] or 's/n'}, "
            f"{fmt_br(m['valor'])}) existia em {m['arquivo']} (linha {m['linha']}) e não consta em {nome_novo}. "
            "Mantido no fechamento (nada some em silêncio); pode ser estorno ou correção do banco. "
            "Helena confirma se continua valendo.",
            [(f"aurora@{m['arquivo']}", m["ref"], m["valor"], "lançamento que sumiu da nova versão (mantido)")]))
    for m in res["movs"]:
        d["rastro_arquivos"][f"aurora@{m['arquivo']}"] = m["arquivo"]
    d["versoes"]["aurora"] = _entrada_versao("aurora", [a for a, _, _ in versoes], res["usados"], res["novas"], 0,
                                             len(res["sumiram"]), res["sobreposta"])
    d["pend_versoes"] += pend
    return res["movs"], res["meta"], (nome_novo if res["sobreposta"] else res["usados"][0])


def _versao_referencia(d, fonte, arqs):
    """Planilha de taxas / plano de contas em 2+ versões: vale a mais recente (ordem de nome); a outra é ignorada."""
    usado, ignorados = arqs[-1].name, [p.name for p in arqs[:-1]]
    rotulo = config.ROTULO_FONTE[fonte]
    d["versoes"][fonte] = _entrada_versao(fonte, [p.name for p in arqs], [usado])
    d["pend_versoes"].append(_pend_versao(
        config.TIPO_VERSAO_SUBSTITUIDA, usado, None,
        f"Há {len(arqs)} versões do arquivo '{rotulo}' na entrada. Usada a mais recente: {usado}. "
        f"Ignorada(s): {', '.join(ignorados)}. A ordem é pelo nome/versão do arquivo (a data de modificação não é "
        "confiável em upload). Informativa: se a ignorada é que deveria valer, remover ou renomear os arquivos.", []))


def _versao_orbe(d, arqs):
    nomes = [p.name for p in arqs]
    subst = sum(1 for x in d["orbe_duplicadas"] if not x["identica"] and x["original"]["arquivo"] == nomes[-1])
    novas = sum(1 for l in d["orbe"] if l["arquivo"] == nomes[-1]) - subst
    conf = sum(1 for x in d["orbe_duplicadas"] if not x["identica"])
    d["versoes"]["orbe"] = _entrada_versao("orbe", nomes, nomes, novas, conf, 0)


def ler_tudo_parcial(pasta):
    """Como ler_tudo, mas tolerante: fonte ausente ou com layout inesperado NÃO interrompe.
    Devolve o mesmo dict (fonte indisponível = vazia) mais:
      'ok'        : conjunto das fontes lidas com sucesso;
      'bloqueios' : [{'fonte', 'tipo' (FONTE_AUSENTE | LAYOUT_INESPERADO), 'arquivo', 'detalhe'}];
      'versoes'   : por fonte com 2+ arquivos, o resumo da consolidação (arquivos, usados, novas, conflitos, texto);
      'pend_versoes' : pendências da consolidação (VERSAO_*, LANCAMENTO_SUMIU_NA_NOVA_VERSAO);
      'rastro_arquivos' : chave de rastro 'fonte@arquivo' -> nome mostrado em rastro.csv;
      'origem_ponte' : {Nº Operação: e-mail} (só com 2+ relatórios da Ponte).
    Com 1 arquivo por fonte nada disso muda o resultado."""
    achados, ausentes = descobrir_versoes(pasta)
    ok, bloqueios = set(), []
    d = {"arquivos": {}, "arquivos_orbe": [], "orbe": [], "orbe_duplicadas": [], "contratos": [],
         "ponte": [], "meta_ponte": {"anexo": None, "assunto": None, "data_repasse": None},
         "plano": {}, "aurora": [], "meta_aurora": {}, "meridiano": [], "meta_meridiano": {},
         "versoes": {}, "pend_versoes": [], "rastro_arquivos": {}, "origem_ponte": {}}

    def bloquear(fonte, tipo, arquivo, detalhe):
        bloqueios.append({"fonte": fonte, "tipo": tipo, "arquivo": arquivo, "detalhe": detalhe})

    def layout(fonte, caminho, prob):
        faltam = ", ".join(prob["faltantes"])
        achou = ", ".join(prob["encontradas"]) or "nenhuma"
        bloquear(fonte, "LAYOUT_INESPERADO", Path(caminho).name,
                 f"layout inesperado em {Path(caminho).name}: faltam {faltam}; encontradas: {achou}")

    def layout_leitura(fonte, caminho, e):
        layout(fonte, caminho, {"faltantes": [f"conteúdo não lido ({type(e).__name__}: {e})"], "encontradas": []})

    for fonte in ausentes:
        bloquear(fonte, "FONTE_AUSENTE", config.PADROES_ARQUIVOS[fonte], "arquivo não encontrado")

    def ler(fonte, caminho, leitor):
        prob = validar_layout(fonte, caminho)
        if prob:
            layout(fonte, caminho, prob)
            return None
        try:
            return leitor(caminho)
        except Exception as e:  # noqa: BLE001
            layout_leitura(fonte, caminho, e)
            return None

    if "orbe" in achados:
        validos = []
        for c in achados["orbe"]:
            if ler("orbe", c, ler_orbe) is not None:
                validos.append(c)
        # S53: qualquer arquivo Orbe ruim bloqueia o checkout inteiro (calcular só com parte dá nota/saldo errados)
        if validos and len(validos) == len(achados["orbe"]):
            d["orbe"], d["orbe_duplicadas"] = ler_orbe_varios(validos)
            d["arquivos_orbe"] = [p.name for p in validos]
            d["arquivos"]["orbe"] = ", ".join(d["arquivos_orbe"])
            ok.add("orbe")
            if len(validos) > 1:
                _versao_orbe(d, validos)
    consolidadores = {"ponte": _consolidar_ponte, "meridiano": _consolidar_ofx}
    for fonte, leitor, campos in (("taxas", ler_taxas, ("contratos",)),
                                  ("ponte", ler_ponte, ("ponte", "meta_ponte")),
                                  ("plano", ler_plano, ("plano",)),
                                  ("aurora", ler_aurora, ("aurora", "meta_aurora")),
                                  ("meridiano", ler_meridiano, ("meridiano", "meta_meridiano"))):
        arqs = achados.get(fonte)
        if not arqs:
            continue
        if fonte in config.FONTES_REFERENCIA or len(arqs) == 1:
            usado = arqs[-1]                       # referência: só a versão mais recente é lida
            r = ler(fonte, usado, leitor)
            if r is None:
                continue
            nome = usado.name
            if len(arqs) > 1:
                _versao_referencia(d, fonte, arqs)
        else:
            lidos, falhou = [], False
            for c in arqs:                          # toda versão precisa ter o layout esperado (como na Orbe)
                rc = ler(fonte, c, leitor)
                if rc is None:
                    falhou = True
                else:
                    lidos.append((c, rc))
            if falhou:
                continue
            if fonte == "aurora":
                movs, meta, nome = _consolidar_aurora_fonte(d, lidos)
                r = (movs, meta)
            else:
                r = consolidadores[fonte](d, lidos)
                nome = arqs[-1].name
        for campo, valor in zip(campos, r if len(campos) > 1 else (r,)):
            d[campo] = valor
        d["arquivos"][fonte] = nome
        ok.add(fonte)
    d["ok"], d["bloqueios"] = ok, bloqueios
    return d
