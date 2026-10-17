"""Versões do mesmo arquivo (reenvio): 2+ arquivos da mesma fonte na pasta de entrada são consolidados por chave
natural, sem duplicar e sem apagar linha em silêncio. Tudo em tmp_path; os arquivos da raiz só são copiados."""
import csv
import email
import io
import shutil
from decimal import Decimal as D
from email import policy

import openpyxl
import pytest

from src import fechamento, leitores

CSVS = ["comissoes_por_lojista.csv", "divergencias_taxa.csv", "notas_a_emitir.csv",
        "lancamentos_classificados.csv", "pendencias.csv", "rastro.csv"]
FIM_AURORA = "\n;;;;\n;SALDO FINAL;;;-1.920.088,60\n"


# ---------------------------------------------------------------- ferramentas
def montar(raiz, destino, sem=()):
    """Copia os arquivos da raiz (menos as fontes em `sem`) para a pasta de entrada."""
    destino.mkdir(parents=True, exist_ok=True)
    for k, v in leitores.descobrir_arquivos(raiz).items():
        if k in sem:
            continue
        for p in (v if isinstance(v, list) else [v]):
            shutil.copy(p, destino / p.name)
    return destino


def rodar(ent, tmp_path, investigar=False):
    saida = tmp_path / "saida"
    return fechamento.executar(ent, saida, investigar=investigar), saida


def ler_csv(saida, nome):
    with open(saida / nome, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f, delimiter=";"))


def pend(r, tipo):
    return [p for p in r["pendencias"] if p["tipo"] == tipo]


def br(v):
    s = f"{abs(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return ("-" if v < 0 else "") + s


def aurora_texto(raiz):
    return leitores.descobrir_arquivos(raiz)["aurora"].read_text(encoding="utf-8-sig")


def aurora_com(texto, *novas):
    """Acrescenta lançamentos no fim do Aurora e corrige o saldo final. novas = (descrição, valor Decimal)."""
    assert FIM_AURORA in texto
    linhas = "".join(f"\n30/09/2026;{d};;{br(v)};0,00" for d, v in novas)
    saldo = D("-1920088.60") + sum((v for _, v in novas), D("0"))
    return texto.replace(FIM_AURORA, f"{linhas}\n;;;;\n;SALDO FINAL;;;{br(saldo)}\n")


def aurora_arquivos(ent, **versoes):
    for nome, texto in versoes.items():
        (ent / f"{nome}.csv").write_text(texto, encoding="utf-8", newline="")


def ofx_texto(raiz):
    return leitores.descobrir_arquivos(raiz)["meridiano"].read_bytes().decode("cp1252")


def ofx_com(texto, dtserver, fitid=None, valor="-12.34", trocar=None):
    """Troca o DTSERVER; opcionalmente acrescenta um movimento ou troca o valor de um FITID."""
    t = texto.replace("<DTSERVER>20261001063000", f"<DTSERVER>{dtserver}")
    if fitid:
        mov = ("<STMTTRN>\r\n<TRNTYPE>DEBIT\r\n<DTPOSTED>20260930120000[-3:BRT]\r\n"
               f"<TRNAMT>{valor}\r\n<FITID>{fitid}\r\n<CHECKNUM>1\r\n<MEMO>TARIFA TESTE\r\n</STMTTRN>\r\n")
        t = t.replace("</BANKTRANLIST>", mov + "</BANKTRANLIST>")
    if trocar:
        alvo, novo = trocar
        i = t.index(f"<FITID>{alvo}")
        j = t.rindex("<TRNAMT>", 0, i)
        k = t.index("\r\n", j)
        t = t[:j] + f"<TRNAMT>{novo}" + t[k:]
    return t


def ofx_arquivos(ent, **versoes):
    for nome, texto in versoes.items():
        (ent / f"{nome}.ofx").write_bytes(texto.encode("cp1252"))


def eml_variante(origem, destino, data=None, ajustar=None):
    """Copia o e-mail da Ponte, com outro cabeçalho Date e/ou o anexo .xlsx ajustado por ajustar(ws, mapa, ops)."""
    with open(origem, "rb") as f:
        msg = email.message_from_binary_file(f, policy=policy.default)
    if data:
        msg.replace_header("Date", data)
    if ajustar:
        parte = next(p for p in msg.iter_attachments() if (p.get_filename() or "").lower().endswith(".xlsx"))
        nome = parte.get_filename()
        wb = openpyxl.load_workbook(io.BytesIO(parte.get_payload(decode=True)), data_only=True)
        ws = wb[wb.sheetnames[0]] if "Operações" not in wb.sheetnames else wb["Operações"]
        _, _, cab, mapa, _ = leitores._abrir_ponte(origem)
        ajustar(ws, mapa, cab)
        buf = io.BytesIO()
        wb.save(buf)
        parte.set_content(buf.getvalue(), maintype="application",
                          subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=nome)
    destino.write_bytes(msg.as_bytes())
    return destino


# ---------------------------------------------------------------- garantia: 1 arquivo por fonte = saida/
def test_um_arquivo_por_fonte_identico_a_saida_byte_a_byte(raiz, tmp_path):
    ent = montar(raiz, tmp_path / "e")
    r, saida = rodar(ent, tmp_path, investigar=True)
    assert set(r["versoes"]) == {"orbe"}          # setembro já tem Orbe + ajustes tardios; as demais fontes, 1 arquivo
    for p in sorted((raiz / "saida").rglob("*")):
        if p.is_file():
            assert (saida / p.relative_to(raiz / "saida")).read_bytes() == p.read_bytes(), p.name


# ---------------------------------------------------------------- ordem das versões
def test_ordem_pelo_nome_natural_e_duplicata_do_navegador(tmp_path):
    def ordem(*nomes):
        arqs = [tmp_path / n for n in nomes]
        for a in arqs:
            a.write_text("")
        return [p.name for p in leitores.ordenar_versoes("plano", arqs)]

    assert ordem("x_v10.csv", "x_v2.csv", "x_v1.csv") == ["x_v1.csv", "x_v2.csv", "x_v10.csv"]
    assert ordem("x (1).csv", "x.csv") == ["x.csv", "x (1).csv"]


def test_ordem_ofx_pelo_dtserver_e_nao_pelo_nome(raiz, tmp_path):
    ent = tmp_path / "e"
    ent.mkdir()
    base = ofx_texto(raiz)
    ofx_arquivos(ent, a_novo=ofx_com(base, "20261003063000[-3:BRT]"), b_velho=ofx_com(base, "20261001063000[-3:BRT]"))
    assert [p.name for p in leitores.ordenar_versoes("meridiano", ent.glob("*.ofx"))] == ["b_velho.ofx", "a_novo.ofx"]


def test_ordem_ponte_pelo_date_do_email(raiz, tmp_path):
    eml = leitores.descobrir_arquivos(raiz)["ponte"]
    a = eml_variante(eml, tmp_path / "a_novo.eml", data="Mon, 05 Oct 2026 07:15:03 -0300")
    b = eml_variante(eml, tmp_path / "b_velho.eml", data="Fri, 02 Oct 2026 07:15:03 -0300")
    assert [p.name for p in leitores.ordenar_versoes("ponte", [a, b])] == ["b_velho.eml", "a_novo.eml"]


def test_ordem_aurora_pelo_fim_do_periodo(raiz, tmp_path):
    base = aurora_texto(raiz)
    a = tmp_path / "a_novo.csv"
    a.write_text(base, encoding="utf-8")
    b = tmp_path / "b_velho.csv"
    b.write_text(base.replace("01/09/2026 a 30/09/2026", "01/09/2026 a 15/09/2026"), encoding="utf-8")
    assert [p.name for p in leitores.ordenar_versoes("aurora", [a, b])] == ["b_velho.csv", "a_novo.csv"]


# ---------------------------------------------------------------- Aurora
def test_aurora_reenviado_com_um_lancamento_a_mais_usa_o_novo(raiz, tmp_path):
    base = aurora_texto(raiz)
    v2 = aurora_com(base, ("TAR TESTE NOVA", D("-10.00")))
    ent = montar(raiz, tmp_path / "e", sem={"aurora"})
    aurora_arquivos(ent, aurora_v1=base, aurora_v2=v2)
    r, saida = rodar(ent, tmp_path)
    aur = [l for l in r["lancamentos"] if l["banco"] == "AURORA"]
    assert len(aur) == 85
    assert all("@" not in l["ref"] for l in aur)                       # formato do edital: Lnn do arquivo usado
    nova = next(l for l in aur if l["descricao"] == "TAR TESTE NOVA")
    assert nova["ref"] == f"L{v2.split(chr(10)).index('30/09/2026;TAR TESTE NOVA;;-10,00;0,00') + 1}"
    rastro = [x for x in ler_csv(saida, "rastro.csv") if x["id_saida"] == nova["id"]]
    assert [x["arquivo_origem"] for x in rastro] == ["aurora_v2.csv"]
    assert not [p for p in r["pendencias"] if p["tipo"].startswith("VERSAO_") or p["tipo"] in
                ("LANCAMENTO_SUMIU_NA_NOVA_VERSAO", "EXTRATO_NAO_FECHA")]
    v = r["versoes"]["aurora"]
    assert (v["arquivos"], v["usados"], v["novas"], v["conflitos"], v["sumiram"]) == (
        ["aurora_v1.csv", "aurora_v2.csv"], ["aurora_v2.csv"], 1, 0, 0)
    # equivale a enviar só a versão nova
    ent2 = montar(raiz, tmp_path / "e2", sem={"aurora"})
    aurora_arquivos(ent2, aurora_v2=v2)
    _, saida2 = rodar(ent2, tmp_path / "x")
    for nome in CSVS:
        assert (saida / nome).read_bytes() == (saida2 / nome).read_bytes(), nome


def test_aurora_resumo_do_cli_com_duas_versoes(raiz, tmp_path):
    base = aurora_texto(raiz)
    ent = montar(raiz, tmp_path / "e", sem={"aurora"})
    aurora_arquivos(ent, aurora_v1=base, aurora_v2=aurora_com(base, ("TAR TESTE NOVA", D("-10.00"))))
    r, _ = rodar(ent, tmp_path)
    assert "Aurora: 2 versões — usada aurora_v2.csv (+1 lançamento novo)" in fechamento.resumo(r).splitlines()[-1] \
        or "    Aurora: 2 versões — usada aurora_v2.csv (+1 lançamento novo)" in fechamento.resumo(r)


def test_aurora_reenviado_identico_nao_muda_nada(raiz, tmp_path):
    base = aurora_texto(raiz)
    ent = montar(raiz, tmp_path / "e", sem={"aurora"})
    aurora_arquivos(ent, aurora_v1=base, aurora_v2=base)
    r, saida = rodar(ent, tmp_path)
    assert len([l for l in r["lancamentos"] if l["banco"] == "AURORA"]) == 84
    assert r["versoes"]["aurora"]["novas"] == 0 and not any(p["tipo"].startswith("VERSAO_") for p in r["pendencias"])
    for nome in CSVS[:5]:
        assert (saida / nome).read_bytes() == (raiz / "saida" / nome).read_bytes(), nome


def test_aurora_duas_tarifas_iguais_no_mesmo_dia_contam_duas_vezes(raiz, tmp_path):
    base = aurora_texto(raiz)
    v1 = aurora_com(base, ("TAR IGUAL", D("-5.00")), ("TAR IGUAL", D("-5.00")))
    v2 = aurora_com(base, ("TAR IGUAL", D("-5.00")), ("TAR IGUAL", D("-5.00")), ("TAR OUTRA", D("-1.00")))
    ent = montar(raiz, tmp_path / "e", sem={"aurora"})
    aurora_arquivos(ent, aurora_v1=v1, aurora_v2=v2)
    r, _ = rodar(ent, tmp_path)
    aur = [l for l in r["lancamentos"] if l["banco"] == "AURORA"]
    assert len(aur) == 84 + 3
    assert sum(1 for l in aur if l["descricao"] == "TAR IGUAL") == 2            # 2, não 4
    assert not any(p["tipo"].startswith("VERSAO_") for p in r["pendencias"])
    assert r["versoes"]["aurora"]["novas"] == 1


def test_aurora_sobreposicao_parcial_une_e_pergunta(raiz, tmp_path):
    base = aurora_texto(raiz)
    ent = montar(raiz, tmp_path / "e", sem={"aurora"})
    aurora_arquivos(ent, aurora_v1=aurora_com(base, ("TAR SO NA ANTIGA", D("-3.00"))),
                    aurora_v2=aurora_com(base, ("TAR SO NA NOVA", D("-4.00"))))
    r, saida = rodar(ent, tmp_path)
    aur = [l for l in r["lancamentos"] if l["banco"] == "AURORA"]
    assert len(aur) == 84 + 2
    assert all(l["ref"].startswith("L") and l["ref"].split("@")[1] in ("aurora_v1.csv", "aurora_v2.csv") for l in aur)
    assert [p["ref_origem"] for p in pend(r, "VERSAO_SOBREPOSTA")] == ["AURORA"]
    assert "UNIÃO" in pend(r, "VERSAO_SOBREPOSTA")[0]["motivo"] and "fecha" in pend(r, "VERSAO_SOBREPOSTA")[0]["motivo"]
    (sumiu,) = pend(r, "LANCAMENTO_SUMIU_NA_NOVA_VERSAO")
    assert sumiu["valor"] == D("-3.00") and sumiu["ref_origem"].endswith("@aurora_v1.csv")
    assert not pend(r, "EXTRATO_NAO_FECHA")
    arqs = {x["ref_origem"]: x["arquivo_origem"] for x in ler_csv(saida, "rastro.csv") if "@" in x["ref_origem"]}
    assert set(arqs.values()) == {"aurora_v1.csv", "aurora_v2.csv"}
    assert all(ref.endswith("@" + arq) for ref, arq in arqs.items())
    assert "união" in r["versoes"]["aurora"]["texto"]


def test_aurora_linha_que_some_na_nova_versao_e_mantida_com_pendencia(raiz, tmp_path):
    base = aurora_texto(raiz)
    ent = montar(raiz, tmp_path / "e", sem={"aurora"})
    aurora_arquivos(ent, aurora_v1=aurora_com(base, ("TAR ESTORNADA", D("-7.00"))), aurora_v2=base)
    r, saida = rodar(ent, tmp_path)
    aur = [l for l in r["lancamentos"] if l["banco"] == "AURORA"]
    assert len(aur) == 85 and any(l["descricao"] == "TAR ESTORNADA" for l in aur)        # nunca apaga em silêncio
    (sumiu,) = pend(r, "LANCAMENTO_SUMIU_NA_NOVA_VERSAO")
    assert sumiu["valor"] == D("-7.00") and "aurora_v1.csv" in sumiu["motivo"] and "aurora_v2.csv" in sumiu["motivo"]
    assert not pend(r, "VERSAO_SOBREPOSTA") and not pend(r, "EXTRATO_NAO_FECHA")
    assert r["versoes"]["aurora"]["usados"] == ["aurora_v1.csv"] and r["versoes"]["aurora"]["sumiram"] == 1
    rastro = [x for x in ler_csv(saida, "rastro.csv") if x["ref_origem"] == sumiu["ref_origem"]]
    assert {x["arquivo_origem"] for x in rastro} == {"aurora_v1.csv"}


def test_aurora_versao_que_nao_fecha_o_saldo_avisa(raiz, tmp_path):
    base = aurora_texto(raiz)
    quebrada = aurora_com(base, ("TAR NOVA", D("-10.00"))).replace(";SALDO FINAL;;;-1.920.098,60", ";SALDO FINAL;;;-1.000,00")
    ent = montar(raiz, tmp_path / "e", sem={"aurora"})
    aurora_arquivos(ent, aurora_v1=base, aurora_v2=quebrada)
    r, _ = rodar(ent, tmp_path)
    assert [p["ref_origem"] for p in pend(r, "EXTRATO_NAO_FECHA")] == ["AURORA"]      # versão usada (v2) não fecha


def test_aurora_versao_com_layout_ruim_bloqueia_o_extrato(raiz, tmp_path):
    base = aurora_texto(raiz)
    ent = montar(raiz, tmp_path / "e", sem={"aurora"})
    aurora_arquivos(ent, aurora_v1=base, aurora_v2=base.replace("Valor (R$)", "Montante"))
    r, _ = rodar(ent, tmp_path)
    assert r["parcial"] and r["bloqueado"]["AURORA"]
    assert [p["ref_origem"] for p in pend(r, "LAYOUT_INESPERADO")] == ["aurora_v2.csv"]


# ---------------------------------------------------------------- Meridiano (OFX)
def test_ofx_reenviado_com_fitid_novo_traz_um_a_mais(raiz, tmp_path):
    base = ofx_texto(raiz)
    ent = montar(raiz, tmp_path / "e", sem={"meridiano"})
    ofx_arquivos(ent, meridiano_v1=base, meridiano_v2=ofx_com(base, "20261002063000[-3:BRT]", fitid="MRD-NOVO-1"))
    r, saida = rodar(ent, tmp_path)
    mer = [l for l in r["lancamentos"] if l["banco"] == "MERIDIANO"]
    assert len(mer) == 101 and sum(1 for l in mer if l["ref"] == "MRD-NOVO-1") == 1
    assert not [p for p in r["pendencias"] if p["tipo"].startswith("VERSAO_") or p["tipo"].startswith("LANCAMENTO_SUMIU")]
    v = r["versoes"]["meridiano"]
    assert v["novas"] == 1 and v["usados"] == ["meridiano_v2.ofx"] and v["conflitos"] == 0
    assert {x["arquivo_origem"] for x in ler_csv(saida, "rastro.csv") if x["ref_origem"] == "MRD-NOVO-1"} == {"meridiano_v2.ofx"}


def test_ofx_fitid_igual_com_valor_diferente_vale_o_mais_recente_e_abre_pendencia(raiz, tmp_path):
    base = ofx_texto(raiz)
    ent = montar(raiz, tmp_path / "e", sem={"meridiano"})
    ofx_arquivos(ent, meridiano_v1=base,
                 meridiano_v2=ofx_com(base, "20261002063000[-3:BRT]", trocar=("MRD2026090003", "207215.78")))
    r, _ = rodar(ent, tmp_path)
    mer = [l for l in r["lancamentos"] if l["banco"] == "MERIDIANO"]
    assert len(mer) == 100
    assert next(l for l in mer if l["ref"] == "MRD2026090003")["valor"] == D("207215.78")     # a mais recente vale
    (p,) = pend(r, "VERSAO_CONFLITANTE")
    assert p["ref_origem"] == "MRD2026090003" and p["quem_decide"] == "Helena (Financeiro)"
    assert "207.215,87" in p["motivo"] and "207.215,78" in p["motivo"]                     # os dois valores
    assert r["versoes"]["meridiano"]["conflitos"] == 1


def test_ofx_movimento_que_some_na_nova_e_mantido_com_pendencia(raiz, tmp_path):
    base = ofx_texto(raiz)
    ent = montar(raiz, tmp_path / "e", sem={"meridiano"})
    ofx_arquivos(ent, meridiano_v1=ofx_com(base, "20261001063000[-3:BRT]", fitid="MRD-SUMIU", valor="-8.00"),
                 meridiano_v2=ofx_com(base, "20261002063000[-3:BRT]"))
    r, _ = rodar(ent, tmp_path)
    assert len([l for l in r["lancamentos"] if l["banco"] == "MERIDIANO"]) == 101
    (p,) = pend(r, "LANCAMENTO_SUMIU_NA_NOVA_VERSAO")
    assert p["ref_origem"] == "MRD-SUMIU" and p["valor"] == D("-8.00")


# ---------------------------------------------------------------- Ponte
def test_ponte_reenviada_identica_nao_muda_nada(raiz, tmp_path):
    eml = leitores.descobrir_arquivos(raiz)["ponte"]
    ent = montar(raiz, tmp_path / "e", sem={"ponte"})
    shutil.copy(eml, ent / "ponte_a.eml")
    shutil.copy(eml, ent / "ponte_b.eml")
    r, saida = rodar(ent, tmp_path)
    rem = [c for c in r["comissoes"] if c["linha_negocio"] == "REMESSA"]
    assert sum((c["comissao"] for c in rem), D("0")) == D("39543.45")
    assert not [p for p in r["pendencias"] if p["tipo"].startswith("VERSAO_") or p["tipo"].startswith("LANCAMENTO_SUMIU")]
    assert len(pend(r, "OPERACAO_DUPLICADA")) == 1                    # a anomalia DENTRO do relatório segue como era
    for nome in CSVS[:5]:
        assert (saida / nome).read_bytes() == (raiz / "saida" / nome).read_bytes(), nome
    rastro = ler_csv(saida, "rastro.csv")
    assert {x["arquivo_origem"].split("#")[0] for x in rastro if x["arquivo_origem"].endswith(".xlsx") or "#" in x["arquivo_origem"]} \
        == {"ponte_b.eml"}
    v = r["versoes"]["ponte"]
    assert v["usados"] == ["ponte_b.eml"] and v["novas"] == 0 and v["conflitos"] == 0


def test_ponte_valor_diferente_na_nova_vale_o_mais_recente_e_abre_pendencia(raiz, tmp_path):
    eml = leitores.descobrir_arquivos(raiz)["ponte"]
    ops, _ = leitores.ler_ponte(eml)
    alvo = next(o for o in ops if o["spread_brl"] > 0)

    def ajuste(ws, mapa, cab):
        c = ws.cell(alvo["linha_excel"], mapa["spread_brl"])
        c.value = float(alvo["spread_brl"]) + 1

    ent = montar(raiz, tmp_path / "e", sem={"ponte"})
    shutil.copy(eml, ent / "ponte_v1.eml")
    eml_variante(eml, ent / "ponte_v2.eml", data="Mon, 05 Oct 2026 07:15:03 -0300", ajustar=ajuste)
    r, _ = rodar(ent, tmp_path)
    (p,) = pend(r, "VERSAO_CONFLITANTE")
    assert p["ref_origem"] == alvo["op"] and p["valor"] == alvo["spread_brl"] + 1
    assert "spread_brl" in p["motivo"]
    assert next(o for o in r["remessa"] if o["op"] == alvo["op"])["spread_brl"] == alvo["spread_brl"] + 1
    assert r["versoes"]["ponte"]["conflitos"] == 1


def test_ponte_operacao_que_some_na_nova_e_mantida_com_pendencia(raiz, tmp_path):
    eml = leitores.descobrir_arquivos(raiz)["ponte"]
    ops, _ = leitores.ler_ponte(eml)
    alvo = next(o for o in ops if o["spread_brl"] > 0 and sum(1 for x in ops if x["op"] == o["op"]) == 1)

    def apagar(ws, mapa, cab):
        for c in range(1, ws.max_column + 1):
            ws.cell(alvo["linha_excel"], c).value = None

    ent = montar(raiz, tmp_path / "e", sem={"ponte"})
    shutil.copy(eml, ent / "ponte_v1.eml")
    eml_variante(eml, ent / "ponte_v2.eml", data="Mon, 05 Oct 2026 07:15:03 -0300", ajustar=apagar)
    r, _ = rodar(ent, tmp_path)
    (p,) = pend(r, "LANCAMENTO_SUMIU_NA_NOVA_VERSAO")
    assert p["ref_origem"] == alvo["op"]
    assert any(o["op"] == alvo["op"] for o in r["remessa"])                       # mantida
    assert sum((c["comissao"] for c in r["comissoes"] if c["linha_negocio"] == "REMESSA"), D("0")) == D("39543.45")


# ---------------------------------------------------------------- fontes de referência
@pytest.mark.parametrize("fonte,padrao,novo", [("taxas", "*taxas*.xlsx", "taxas_v2.xlsx"),
                                              ("plano", "*plano_de_contas*.csv", "plano_de_contas_v2.csv")])
def test_referencia_em_duas_versoes_usa_a_mais_recente(raiz, tmp_path, fonte, padrao, novo):
    orig = leitores.descobrir_arquivos(raiz)[fonte]
    ent = montar(raiz, tmp_path / "e", sem={fonte})
    velho = ent / orig.name.replace(".", "_v1.", 1)
    velho.write_bytes(b"conteudo antigo que nao precisa nem abrir")        # prova que a versão ignorada não é lida
    shutil.copy(orig, ent / novo)
    r, _ = rodar(ent, tmp_path)
    assert not r["parcial"]
    assert r["dados"]["arquivos"][fonte] == novo
    (p,) = pend(r, "VERSAO_SUBSTITUIDA")
    assert novo in p["motivo"] and velho.name in p["motivo"] and p["valor"] is None
    v = r["versoes"][fonte]
    assert v["usados"] == [novo] and v["ignorados"] == [velho.name]
    assert f"usada {novo} (ignorada {velho.name})" in v["texto"]
    # os números do mês não mudam
    assert sum((c["comissao"] for c in r["comissoes"] if c["linha_negocio"] == "REMESSA"), D("0")) == D("39543.45")


# ---------------------------------------------------------------- Orbe (já existia) no resumo de versões
def test_orbe_reenvio_identico_aparece_no_resumo_e_nao_soma_em_dobro(raiz, tmp_path):
    ent = montar(raiz, tmp_path / "e")                                         # 01_orbe + 06_orbe_ajustes (setembro)
    shutil.copy(ent / "01_orbe_relatorio_comissoes_set2026.csv", ent / "00_orbe_copia.csv")   # reenvio idêntico do 01
    r, saida = rodar(ent, tmp_path)
    v = r["versoes"]["orbe"]
    assert v["arquivos"] == ["00_orbe_copia.csv", "01_orbe_relatorio_comissoes_set2026.csv",
                             "06_orbe_ajustes_tardios_set2026.csv"]
    assert "Orbe: 3 arquivos" in fechamento.resumo(r)
    for nome in ("comissoes_por_lojista.csv", "notas_a_emitir.csv", "divergencias_taxa.csv"):
        assert (saida / nome).read_bytes() == (raiz / "saida" / nome).read_bytes(), nome


def test_resumo_lista_so_as_fontes_com_mais_de_um_arquivo(raiz, tmp_path):
    base = aurora_texto(raiz)
    ent = montar(raiz, tmp_path / "e", sem={"aurora", "meridiano"})
    aurora_arquivos(ent, aurora_v1=base, aurora_v2=aurora_com(base, ("TAR TESTE NOVA", D("-10.00"))))
    ofx_arquivos(ent, meridiano_v1=ofx_texto(raiz))
    r, _ = rodar(ent, tmp_path)
    linhas = fechamento.resumo(r).splitlines()
    assert "  Versões de arquivos (reenvios consolidados, sem duplicar):" in linhas
    assert "    Aurora: 2 versões — usada aurora_v2.csv (+1 lançamento novo)" in linhas
    assert not any(l.strip().startswith("Meridiano:") for l in linhas)


def test_sobreposicao_aurora_com_lancamento_corrigido_lista_candidatos(raiz, tmp_path):
    base = aurora_texto(raiz)
    ent = montar(raiz, tmp_path / "e", sem={"aurora"})
    aurora_arquivos(ent, aurora_v1=aurora_com(base, ("TAR CORRIGIDA", D("-3.00")), ("TAR SO NA ANTIGA", D("-1.00"))),
                    aurora_v2=aurora_com(base, ("TAR CORRIGIDA", D("-4.00"))))
    r, _ = rodar(ent, tmp_path)
    (p,) = pend(r, "VERSAO_SOBREPOSTA")
    assert "remova a versão antiga na aba Fechamento (botão 'remover desta pasta')" in p["motivo"]
    assert "os dois lançamentos estão contados" in p["motivo"]
    assert "'TAR CORRIGIDA' de 30/09/2026: R$ -3,00 em aurora_v1.csv -> R$ -4,00 em aurora_v2.csv" in p["motivo"]
    assert "TAR SO NA ANTIGA" not in p["motivo"].split("Candidatos")[1]


def test_divergencia_liquida_negativa_nao_manda_devolver(raiz, tmp_path):
    from src import cadastro
    cad = tmp_path / "cadastro"
    shutil.copytree(raiz / "cadastro", cad)
    for faixa in ("DEBITO", "CREDITO A VISTA", "PARC_2_6", "PARC_7_12", "PIX"):
        cadastro.adicionar(cad, {"tipo": "TAXA_ESPECIAL", "cnpj": "95927868000148", "lojista": "Rota Global",
                                 "linha_negocio": "CHECKOUT", "campo": faixa, "valor": "9.00",
                                 "vigencia_inicio": "2026-09-01", "aprovado_por": "Comercial", "fonte": "teste"}, "ana")
    ent = montar(raiz, tmp_path / "e")
    r = fechamento.executar(ent, tmp_path / "saida", cadastro=cad)
    (p,) = [x for x in r["pendencias"] if x["tipo"] == "DIVERGENCIA_TAXA_CHECKOUT" and x["valor"] < 0]
    assert "a menos que o contrato; não há devolução ao lojista; tratar caso a caso (R50/S24)" in p["motivo"]
    assert "Devolver" not in p["motivo"] and "a devolver" not in p["motivo"]
    g = next(g for g in r["investigacao"]["grupos"] if g["lojista"].startswith("Rota Global") and float(g["valor_devolver_lojista_brl"]) < 0)
    assert "Devolver" not in g["acao_regra"]["texto"] and "a menos que o contrato" in g["acao_regra"]["texto"]
    assert not any("A devolver" in f["texto"] for f in g["fatos"] if "líquido" in f["texto"])
